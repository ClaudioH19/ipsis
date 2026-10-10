"""Construcción explícita del corpus sintético; no se regenera al evaluarlo."""
from __future__ import annotations

import argparse
import json
from pathlib import Path
import textwrap

TEMAS = {
    "rrhh": {
        "titulo": "Procedimiento de teletrabajo", "extension": ".md",
        "hechos": ["Las solicitudes de teletrabajo se presentan con 12 días hábiles de anticipación.",
                   "La jefatura de Personas autoriza las solicitudes de teletrabajo."],
        "preguntas": ["¿Con cuánta anticipación se presentan las solicitudes de teletrabajo?",
                      "¿Quién autoriza las solicitudes de teletrabajo?", "¿Cuántos días antes hay que pedir teletrabajo?"],
        "literales": ["12 días hábiles", "jefatura de Personas", "12 días hábiles"],
        "trampa": "¿Cuánto dinero se paga por cada día de teletrabajo?",
        "premisa": "¿Por qué las solicitudes de teletrabajo se presentan con 3 días de anticipación?",
    },
    "compras": {
        "titulo": "Compras de equipos informáticos", "extension": ".pdf",
        "hechos": ["Las cotizaciones para comprar equipos informáticos tienen una vigencia de 18 días corridos.",
                   "Para iniciar una compra de equipos informáticos, Abastecimiento revisa las cotizaciones."],
        "preguntas": ["¿Cuál es la vigencia de las cotizaciones para comprar equipos informáticos?",
                      "¿Quién revisa las cotizaciones para iniciar una compra de equipos informáticos?",
                      "¿Cuánto duran las cotizaciones de equipos informáticos?"],
        "literales": ["18 días corridos", "Abastecimiento", "18 días corridos"],
        "trampa": "¿Cuál es el precio máximo para comprar equipos informáticos?",
        "premisa": "¿Por qué las cotizaciones de equipos informáticos tienen una vigencia de 5 días?",
    },
    "tecnico": {
        "titulo": "Mantenimiento de la bomba de recirculación", "extension": ".docx",
        "hechos": ["El filtro de la bomba de recirculación se cambia cada 45 días.",
                   "Antes de cambiar el filtro de la bomba de recirculación, se debe desconectar la alimentación eléctrica."],
        "preguntas": ["¿Cada cuánto se cambia el filtro de la bomba de recirculación?",
                      "¿Qué se debe hacer antes de cambiar el filtro de la bomba de recirculación?",
                      "¿Con qué frecuencia hay que reemplazar el filtro de la bomba de recirculación?"],
        "literales": ["45 días", "desconectar la alimentación eléctrica", "45 días"],
        "trampa": "¿Cuánto cuesta cambiar el filtro de la bomba de recirculación?",
        "premisa": "¿Por qué el filtro de la bomba de recirculación se cambia cada 90 días?",
    },
    "logistica": {
        "titulo": "Conservación y despacho de paquetes refrigerados", "extension": ".txt",
        "hechos": ["Los paquetes refrigerados se conservan entre 2 y 6 grados Celsius.",
                   "El despacho de paquetes refrigerados se confirma después de verificar el embalaje."],
        "preguntas": ["¿A qué temperatura se conservan los paquetes refrigerados?",
                      "¿Cuándo se confirma el despacho de paquetes refrigerados?",
                      "¿En qué rango de temperatura hay que mantener los paquetes refrigerados?"],
        "literales": ["entre 2 y 6 grados Celsius", "después de verificar el embalaje", "entre 2 y 6 grados Celsius"],
        "trampa": "¿Cuánto cuesta el despacho de paquetes refrigerados?",
        "premisa": "¿Por qué los paquetes refrigerados se conservan a 15 grados Celsius?",
    },
}
CALIDADES = ("bueno", "medio", "malo")
VERSION = "1.1.0"
AVISO = "Material ficticio para pruebas de software. No describe normas ni productos reales."


def lineas(tema, calidad):
    titulo, hechos = tema["titulo"], tema["hechos"]
    distractor = "Las etiquetas internas se revisan cada 8 días."
    if calidad == "bueno":
        return [titulo, AVISO, "Procedimiento", hechos[0], "Responsabilidades y condiciones", hechos[1],
                "Control documental", distractor]
    if calidad == "medio":
        return [titulo, AVISO, "INFORMACIÓN GENERAL", hechos[0], hechos[1], distractor,
                "Copia interna. Página 1. FIN DEL PROCEDIMIENTO."]
    # Orden irregular, encabezados repetidos, párrafos pegados, ruido y cierre truncado.
    # Los hechos y sus unidades permanecen íntegros: la prueba degrada formato, no verdad.
    return ["COPIA INTERNA / PÁGINA 1", titulo, AVISO,
            "ANEXO - CONTROL / " + distractor, "COPIA INTERNA / PÁGINA 1",
            hechos[1] + "     Referencia: archivo interno.", "----- continuación -----",
            hechos[0], "COPIA INTERNA / PÁGINA 2", "Observaciones adicionales: el documento se",
            "FIN / COPIA INTERNA"]


def pdf(ruta, textos, calidad):
    from pypdf import PdfWriter
    from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject

    escritor = PdfWriter()
    pagina = escritor.add_blank_page(width=595, height=842)
    fuente = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"),
        NameObject("/BaseFont"): NameObject("/Helvetica"), NameObject("/Encoding"): NameObject("/WinAnsiEncoding")})
    pagina[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): fuente})})
    instrucciones = [b"BT /F1 11 Tf 50 790 Td 15 TL"]
    for texto in textos:
        for linea in textwrap.wrap(texto, width=83):
            contenido = linea.encode("cp1252").replace(b"\\", b"\\\\").replace(b"(", b"\\(").replace(b")", b"\\)")
            instrucciones.append(b"(" + contenido + b") Tj T*")
        instrucciones.append(b"T*")
    instrucciones.append(b"ET")
    stream = DecodedStreamObject()
    stream.set_data(b"\n".join(instrucciones))
    pagina[NameObject("/Contents")] = escritor._add_object(stream)
    if calidad == "bueno":
        escritor.add_outline_item(textos[0], 0)
    with ruta.open("wb") as archivo:
        escritor.write(archivo)


def documento(ruta, tema, calidad):
    textos = lineas(tema, calidad)
    if ruta.suffix == ".pdf":
        pdf(ruta, textos, calidad)
    elif ruta.suffix == ".docx":
        from docx import Document

        doc = Document()
        for i, texto in enumerate(textos):
            if calidad == "bueno" and i in (0, 2, 4, 6):
                doc.add_heading(texto, level=1 if i == 0 else 2)
            elif calidad == "medio" and i == 0:
                doc.add_paragraph().add_run(texto).bold = True
            else:
                doc.add_paragraph(texto)
        doc.save(ruta)
    else:
        if ruta.suffix == ".md" and calidad == "bueno":
            textos = [("# " if i == 0 else "## ") + t if i in (0, 2, 4, 6) else t for i, t in enumerate(textos)]
        separador = "\n" if calidad == "malo" else "\n\n"
        ruta.write_text(separador.join(textos) + "\n", encoding="utf-8", newline="\n")


def construir(destino):
    from ipsis.evaluacion.banco import ArchivoCorpus, hash_archivos_corpus
    from ipsis.ingesta.parsear import parsear_documento

    for calidad in CALIDADES:
        carpeta = destino / calidad
        corpus = carpeta / "corpus"
        corpus.mkdir(parents=True, exist_ok=True)
        archivos, casos = [], []
        for nombre, tema in TEMAS.items():
            ruta = corpus / (nombre + tema["extension"])
            documento(ruta, tema, calidad)
            archivos.append(ArchivoCorpus(ruta=ruta.name, doc_id=nombre))
            texto = parsear_documento(str(ruta), doc_id=nombre, doc_titulo=nombre, version="1").texto_completo
            for i, (pregunta, literal) in enumerate(zip(tema["preguntas"], tema["literales"])):
                inicio = texto.index(literal)
                # Los offsets son un golden congelado. No se recalculan en general.py.
                casos.append({"id": f"{nombre}-{i+1}", "dataset_version": VERSION, "split": "dev" if i == 0 else "test",
                    "categoria": "parafrasis" if i == 2 else "exacta", "pregunta": pregunta,
                    "estado_esperado": "respuesta_directa", "respondible": True, "respuesta_literal": literal,
                    "evidencias": [{"doc_id": nombre, "offset_inicio": inicio, "offset_fin": inicio+len(literal), "texto": literal}],
                    "notas": f"tema={nombre};formato={ruta.suffix[1:]};calidad={calidad}"})
            for i, (tipo, pregunta) in enumerate((("sin_respuesta", tema["trampa"]), ("premisa_falsa", tema["premisa"]),
                                                ("ambigua", "¿Cuánto tarda?"))):
                caso = {"id": f"{nombre}-neg-{i+1}", "dataset_version": VERSION, "split": "test" if i else "dev",
                        "categoria": tipo, "pregunta": pregunta, "respondible": False,
                        "estado_esperado": "necesita_aclaracion" if tipo == "ambigua" else "sin_respuesta",
                        "notas": f"tema={nombre};formato={ruta.suffix[1:]};calidad={calidad}"}
                if tipo == "sin_respuesta":
                    caso.update(subtipo="trampa_mismo_tema", par=f"{nombre}-1")
                casos.append(caso)
        casos += [{"id": "saludo", "dataset_version": VERSION, "split": "dev", "categoria": "conversacional",
                   "pregunta": "hola", "estado_esperado": "conversacional", "respondible": False},
                  {"id": "fuera-dominio", "dataset_version": VERSION, "split": "test", "categoria": "sin_respuesta",
                   "pregunta": "¿Quién ganó la copa mundial de fútbol?", "estado_esperado": "fuera_dominio", "respondible": False}]
        for id_, pregunta in (("referencia-actor", "¿Quién lo aprueba?"), ("referencia-importe", "¿Cuánto es?")):
            casos.append({"id": id_, "dataset_version": VERSION, "split": "test", "categoria": "ambigua",
                          "pregunta": pregunta, "estado_esperado": "necesita_aclaracion", "respondible": False,
                          "notas": "Regresión conocida: la pregunta independiente no tiene antecedente"})
        (carpeta / "banco.jsonl").write_text("".join(json.dumps(c, ensure_ascii=False)+"\n" for c in casos), encoding="utf-8")
        manifiesto = {"schema_version": 1, "dataset_id": f"general_{calidad}_es", "dataset_version": VERSION,
            "tipo": "sintetico", "estado": "estable", "uso": "regresion", "idioma": "es", "casos": "banco.jsonl",
            "corpus": {"directorio": "corpus", "archivos": [a.model_dump() for a in archivos],
                       "sha256": hash_archivos_corpus(corpus, archivos)},
            "categorias_requeridas": ["exacta", "parafrasis", "sin_respuesta", "premisa_falsa", "ambigua", "conversacional"]}
        (carpeta / "manifiesto.json").write_text(json.dumps(manifiesto, ensure_ascii=False, indent=2)+"\n", encoding="utf-8")
        print(calidad, len(casos), "casos", flush=True)


if __name__ == "__main__":
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("--output", type=Path, required=True, help="Construir en una carpeta nueva para revisar los cambios")
    args = p.parse_args()
    if args.output.exists():
        p.error("La carpeta de salida debe ser nueva; no se sobrescribe el golden existente")
    construir(args.output)
