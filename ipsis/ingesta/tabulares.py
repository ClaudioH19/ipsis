"""Inventarios: registros completos y coordenadas de origen, sin cálculo."""
from __future__ import annotations

import csv
from datetime import date, datetime, time
import json
import re
from pathlib import Path

MAX_CELDAS = 500_000
PARSER_VERSION = "tabular-1"
EXTENSIONES = {".xlsx", ".xls", ".csv", ".tsv"}


def _texto(valor, formato="General"):
    if valor is None:
        return ""
    if isinstance(valor, bool):
        return "TRUE" if valor else "FALSE"
    if isinstance(valor, (date, datetime, time)):
        return valor.isoformat()
    if isinstance(valor, (int, float)):
        if isinstance(valor, float) and valor.is_integer():
            valor = int(valor)
        if isinstance(valor, int) and re.fullmatch(r"0+", formato):
            return ("-" if valor < 0 else "") + str(abs(valor)).zfill(len(formato))
    return str(valor)


def _columna(n):
    salida = ""
    while n:
        n, resto = divmod(n - 1, 26)
        salida = chr(65 + resto) + salida
    return salida


def _referencia(hoja, fila, columna):
    return f"'{hoja.replace(chr(39), chr(39)*2)}'!{_columna(columna)}{fila}"


def _celda_excel(celda, almacenada, hoja, avisos):
    valor = celda.value
    formula = valor if celda.data_type == "f" else None
    if formula:
        if almacenada.value is None:
            avisos.append(f"{hoja}!{celda.coordinate}: fórmula sin resultado guardado; se conserva la fórmula")
        else:
            valor = almacenada.value
            avisos.append(f"{hoja}!{celda.coordinate}: resultado de fórmula guardado; no se recalcula")
    return {"text": _texto(valor, celda.number_format), "locator": _referencia(hoja, celda.row, celda.column),
            "value_type": celda.data_type, "number_format": celda.number_format, "formula": formula}


def _xlsx(ruta, avisos):
    from openpyxl import load_workbook
    from openpyxl.utils.cell import range_boundaries

    libro = load_workbook(ruta, data_only=False, keep_links=False)
    valores = None
    try:
        valores = load_workbook(ruta, data_only=True, read_only=True, keep_links=False)
        for hoja in libro.worksheets:
            if hoja.max_row * hoja.max_column > MAX_CELDAS:
                raise ValueError(f"{hoja.title}: rango de más de {MAX_CELDAS} celdas; divide el inventario")
            if hoja.sheet_state != "visible":
                avisos.append(f"Hoja oculta incluida: {hoja.title}")
            if hoja.merged_cells.ranges:
                avisos.append(f"{hoja.title}: celdas combinadas; no se propagan valores entre registros")
            regiones = [(t.displayName, range_boundaries(t.ref), bool(t.headerRowCount)) for t in hoja.tables.values()]
            if not regiones:
                regiones = [(hoja.title, (1, 1, hoja.max_column, hoja.max_row), False)]
            cubiertas = set()
            for nombre, (a, b, c, d), declarado in regiones:
                if (c-a+1) * (d-b+1) > MAX_CELDAS:
                    raise ValueError(f"{nombre}: rango demasiado grande; divide el inventario")
                filas = []
                cache = valores[hoja.title].iter_rows(min_row=b, max_row=d, min_col=a, max_col=c)
                for fila, guardada in zip(hoja.iter_rows(min_row=b, max_row=d, min_col=a, max_col=c), cache):
                    copias = []
                    for celda, almacenada in zip(fila, guardada):
                        copias.append(_celda_excel(celda, almacenada, hoja.title, avisos))
                        cubiertas.add((celda.row, celda.column))
                    if any(c["text"].strip() for c in copias):
                        filas.append((fila[0].row, copias))
                if filas:
                    sin_header = bool(hoja.tables) and not declarado
                    if sin_header:
                        headers = [{"text": _columna(i), "locator": f"columna {_columna(i)} (etiqueta de coordenada)"} for i in range(a, c+1)]
                        yield hoja.title, nombre, (0, headers), filas, False
                        continue
                    if not declarado:
                        avisos.append(f"{hoja.title}: primera fila no vacía usada como encabezados; revisa la vista")
                    yield hoja.title, nombre, filas[0], filas[1:], declarado
            # No descartar notas ni registros fuera de tablas declaradas.
            if hoja.tables:
                extras = []
                cache = valores[hoja.title].iter_rows(min_row=1, max_row=hoja.max_row, min_col=1, max_col=hoja.max_column)
                for fila, almacenadas in zip(hoja.iter_rows(), cache):
                    celdas = [_celda_excel(c, v, hoja.title, avisos) if (c.row, c.column) not in cubiertas
                               else {"text": "", "locator": _referencia(hoja.title, c.row, c.column)}
                               for c, v in zip(fila, almacenadas)]
                    if any(c["text"].strip() for c in celdas):
                        extras.append((fila[0].row, celdas))
                if extras:
                    headers = [{"text": _columna(i+1), "locator": f"columna {_columna(i+1)} (etiqueta de coordenada)"}
                               for i in range(hoja.max_column)]
                    avisos.append(f"{hoja.title}: filas fuera de tablas conservadas por coordenadas de columna")
                    yield hoja.title, "fuera de tablas", (0, headers), extras, False
    finally:
        libro.close()
        if valores is not None:
            valores.close()


def _xls(ruta, avisos):
    import xlrd

    libro = xlrd.open_workbook(str(ruta), formatting_info=True, on_demand=True)
    avisos.append("XLS: se recuperan valores guardados; no se recalculan fórmulas")
    try:
        for hoja in libro.sheets():
            if hoja.nrows * hoja.ncols > MAX_CELDAS:
                raise ValueError(f"{hoja.name}: rango demasiado grande; divide el inventario")
            filas = []
            for y in range(hoja.nrows):
                celdas = []
                for x in range(hoja.ncols):
                    celda = hoja.cell(y, x)
                    valor = celda.value
                    if celda.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                        valor = None
                    elif celda.ctype == xlrd.XL_CELL_DATE:
                        valor = xlrd.xldate_as_datetime(valor, libro.datemode)
                    elif celda.ctype == xlrd.XL_CELL_BOOLEAN:
                        valor = bool(valor)
                    elif celda.ctype == xlrd.XL_CELL_ERROR:
                        valor = xlrd.error_text_from_code.get(valor, f"error XLS {valor}")
                    formato = libro.format_map[libro.xf_list[hoja.cell_xf_index(y, x)].format_key].format_str
                    celdas.append({"text": _texto(valor, formato), "locator": _referencia(hoja.name, y+1, x+1),
                                   "number_format": formato, "value_type": str(celda.ctype)})
                if any(c["text"].strip() for c in celdas):
                    filas.append((y+1, celdas))
            if filas:
                avisos.append(f"{hoja.name}: primera fila no vacía usada como encabezados; revisa la vista")
                yield hoja.name, hoja.name, filas[0], filas[1:], False
    finally:
        libro.release_resources()


def _csv(ruta, avisos):
    contenido = ruta.read_bytes()
    try:
        texto = contenido.decode("utf-8-sig")
    except UnicodeDecodeError:
        texto = contenido.decode("cp1252")
        avisos.append("CSV leído con codificación Windows-1252")
    import io

    delimitador = "\t" if ruta.suffix.lower() == ".tsv" else None
    if delimitador is None:
        try:
            delimitador = csv.Sniffer().sniff(texto[:32768], delimiters=",;\t|").delimiter
        except csv.Error:
            # Sniffer falla con saltos de línea dentro de celdas entrecomilladas.
            # Comparar registros CSV, no líneas físicas, conserva esas celdas.
            from collections import Counter

            candidatos = []
            for d in (",", ";", "\t", "|"):
                registros = [r for r in csv.reader(io.StringIO(texto[:32768], newline=""), delimiter=d) if r]
                if registros and len(registros[0]) > 1:
                    frecuencias = Counter(map(len, registros))
                    candidatos.append((max(frecuencias.values()) / len(registros), len(registros[0]), d))
            delimitador = max(candidatos)[2] if candidatos else ","
    filas, total = [], 0
    for numero, valores in enumerate(csv.reader(io.StringIO(texto, newline=""), delimiter=delimitador, strict=True), 1):
        total += len(valores)
        if total > MAX_CELDAS:
            raise ValueError("CSV demasiado grande; divide el inventario")
        if any(v.strip() for v in valores):
            filas.append((numero, [{"text": v, "locator": f"fila {numero}, columna {_columna(i+1)}"}
                                   for i, v in enumerate(valores)]))
    if filas:
        avisos.append("CSV/TSV: primera fila no vacía usada como encabezados")
        yield ruta.stem, ruta.stem, filas[0], filas[1:], True


def parsear_tabular(ruta, doc_id, doc_titulo, version):
    from ipsis.ingesta.parsear import Bloque, DocumentoParseado

    ruta = Path(ruta)
    avisos, partes, bloques, offset, registros = [], [], [], 0, 0
    lector = {".xlsx": _xlsx, ".xls": _xls, ".csv": _csv, ".tsv": _csv}[ruta.suffix.lower()]

    def agregar(texto, seccion, metadata, tipo="tabla"):
        nonlocal offset
        inicio = offset
        partes.append(texto + "\n")
        offset += len(texto) + 1
        bloques.append(Bloque(len(bloques), seccion, len(seccion), tipo, texto, inicio, offset-1, metadata))

    for hoja, nombre, (fila_header, headers), filas, declarado in lector(ruta, avisos):
        ancho = max([len(headers), *(len(c) for _, c in filas)])
        if not filas:
            avisos.append(f"{hoja}: sin filas de datos")
            continue
        agregar(hoja if nombre == hoja else f"{hoja} / {nombre}", [hoja, nombre], {}, "titulo")
        for fila, datos in filas:
            registro = json.dumps([hoja, nombre, fila], ensure_ascii=False)
            seccion = [hoja] + ([nombre] if nombre != hoja else []) + [f"Fila {fila}"]
            for es_header, celdas in ((True, headers), (False, datos)):
                copias, textos, posicion = [], [], offset
                for columna in range(ancho):
                    fila_ref = fila_header if es_header else fila
                    localizador = (f"fila {fila_ref}, columna {_columna(columna+1)} (posición omitida)"
                        if ruta.suffix.lower() in {".csv", ".tsv"} else _referencia(hoja, fila_ref, columna+1))
                    c = celdas[columna] if columna < len(celdas) else {"text": "", "locator": localizador}
                    if columna:
                        posicion += 3
                    copias.append({**c, "column": columna, "column_span": 1, "row_span": 1,
                                   "offset_start": posicion, "offset_end": posicion + len(c["text"])})
                    textos.append(c["text"])
                    posicion += len(c["text"])
                agregar(" | ".join(textos), seccion, {"table_id": registro, "record_id": registro,
                    "source_sheet": hoja, "source_row": fila, "header": es_header and declarado,
                    "columns": ancho, "cells": copias, "parser_version": PARSER_VERSION})
            registros += 1
    if not registros:
        raise ValueError("el inventario no contiene filas de datos recuperables")
    return DocumentoParseado(doc_id, doc_titulo, version, str(ruta), "".join(partes), bloques,
                             advertencias=sorted(set(avisos)))


def preparar_tabular(ruta, *, doc_id, doc_titulo, version="1"):
    """La misma extracción en el contrato estructural: una tabla por registro."""
    import hashlib
    from ipsis.ingesta.preparacion import _hash
    from ipsis.nucleo.documentos import CeldaPreparada, DocumentoPreparado, ReferenciaFuente, TablaPreparada
    from ipsis.nucleo.esquemas import ElementRecord

    doc = parsear_tabular(ruta, doc_id, doc_titulo, version)
    elementos, tablas, grupos = [], [], {}
    for b in doc.bloques:
        if b.metadata.get("record_id"):
            grupos.setdefault(b.metadata["record_id"], []).append(b)
        else:
            elementos.append(ElementRecord(element_id=f"{doc_id}__e{b.indice}", tipo=b.tipo_contenido,
                texto=b.texto, offset_inicio=b.offset_inicio, offset_fin=b.offset_fin, seccion_path=b.seccion_path))
    for numero, filas in enumerate(grupos.values()):
        tid = f"{doc_id}__tabla{numero}"
        inicio, fin = filas[0].offset_inicio, filas[-1].offset_fin
        elementos.append(ElementRecord(element_id=tid, tipo="tabla", texto=doc.texto_completo[inicio:fin],
            offset_inicio=inicio, offset_fin=fin, seccion_path=filas[0].seccion_path,
            metadata={"tipo": "tabla", "record_id": filas[0].metadata["record_id"]}))
        ids, celdas = [], []
        for y, b in enumerate(filas):
            fid = f"{tid}__fila{y}"
            ids.append(fid)
            elementos.append(ElementRecord(element_id=fid, tipo="tabla", texto=b.texto,
                offset_inicio=b.offset_inicio, offset_fin=b.offset_fin, seccion_path=b.seccion_path,
                metadata={**b.metadata, "tipo": "fila", "parent_id": tid}))
            for c in b.metadata["cells"]:
                celdas.append(CeldaPreparada(cell_id=f"{fid}__c{c['column']}", fila=y, columna=c["column"],
                    texto=c["text"], referencia=ReferenciaFuente(element_id=fid,
                    offset_inicio=c["offset_start"], offset_fin=c["offset_end"], localizador=c["locator"])))
        tablas.append(TablaPreparada(table_id=tid, filas=len(filas), columnas=filas[0].metadata["columns"],
            celdas=celdas, filas_ids=ids, filas_encabezado=[0], encabezado_explicito=filas[0].metadata["header"],
            referencia=ReferenciaFuente(element_id=tid, offset_inicio=inicio, offset_fin=fin,
                localizador=f"{filas[0].metadata['source_sheet']}, fila {filas[0].metadata['source_row']}")))
    elementos.sort(key=lambda e: (e.offset_inicio, e.metadata.get("tipo") == "fila"))
    return DocumentoPreparado(doc_id=doc_id, titulo=doc_titulo, version=str(version), ruta_origen=str(ruta),
        texto_completo=doc.texto_completo, elementos=elementos, tablas=tablas, parser_version=PARSER_VERSION,
        source_hash="sha256:" + hashlib.sha256(Path(ruta).read_bytes()).hexdigest(),
        structure_hash=_hash([e.model_dump(mode="json") for e in elementos]), advertencias=doc.advertencias)
