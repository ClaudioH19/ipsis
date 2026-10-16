"""Alcance de las unidades y rasgos que distinguen a las hermanas (preparación estructural)."""
from __future__ import annotations

from docx import Document

from ipsis.ingesta.hermanos import agrupar_hermanos, rasgos_distintivos
from ipsis.ingesta.preparacion import preparar_docx
from ipsis.ingesta.unidades import aplicar_senales, unidades_documento


def _documento(tmp_path, filas=3):
    d = Document()
    d.add_heading("Evaluación", 1)
    for titulo, pesos in (("Práctica Profesional I", [20, 40, 40]), ("Práctica Profesional II", [20, 35, 45])):
        d.add_paragraph(titulo)
        t = d.add_table(rows=len(pesos) + 1, cols=2)
        t.cell(0, 0).text, t.cell(0, 1).text = "Instrumento", "Peso"
        for i, v in enumerate(pesos, 1):
            t.cell(i, 0).text, t.cell(i, 1).text = f"Instrumento número {i} del plan", f"{v}%"
    archivo = tmp_path / "m.docx"
    d.save(archivo)
    return preparar_docx(archivo, doc_id="m", doc_titulo="Manual", version="1")


def test_el_alcance_de_una_tabla_es_el_parrafo_corto_que_la_precede(tmp_path):
    hijos, padres = unidades_documento(_documento(tmp_path))
    tablas = [p for p in padres.values() if p.metadata["table_id"]]
    assert sorted(p.metadata["alcance"] for p in tablas) == ["Práctica Profesional I", "Práctica Profesional II"]
    assert {h.metadata["alcance"] for h in hijos if h.metadata["table_id"]} == {"Práctica Profesional I", "Práctica Profesional II"}


def test_un_texto_largo_no_es_alcance(tmp_path):
    d = Document()
    d.add_heading("Plazos", 1)
    d.add_paragraph("Este párrafo es una oración completa, demasiado larga para ser una etiqueta de alcance válida.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text, t.cell(1, 0).text, t.cell(1, 1).text = "A", "B", "x", "y"
    archivo = tmp_path / "p.docx"
    d.save(archivo)
    _, padres = unidades_documento(preparar_docx(archivo, doc_id="p", doc_titulo="P", version="1"))
    assert all("alcance" not in p.metadata for p in padres.values())


def test_aplicar_senales_agrega_el_alcance_solo_si_se_pide_y_no_lo_repite(tmp_path):
    documento = _documento(tmp_path)
    hijos, padres = unidades_documento(documento, max_palabras=6)   # tablas partidas: el 2.º tramo pierde la etiqueta
    sin = {h.chunk_id: h.texto_embedding for h in hijos}
    aplicar_senales(hijos, padres, {"alcance": False})
    assert {h.chunk_id: h.texto_embedding for h in hijos} == sin
    aplicar_senales(hijos, padres, {"alcance": True})
    for h in (h for h in hijos if h.metadata.get("alcance")):
        alcance = h.metadata["alcance"]
        assert h.texto_embedding.casefold().count(alcance.casefold()) >= 1
        assert h.texto_embedding.endswith(h.texto_fuente)
        assert h.texto_embedding.count("Alcance: ") <= 1
    con_alcance = [h for h in hijos if h.metadata.get("alcance")]
    assert con_alcance and all(h.metadata["alcance"].casefold() in h.texto_embedding.casefold() for h in con_alcance)


def test_rasgos_separan_a_las_hermanas_y_son_deterministas():
    textos = {
        "p1": "Práctica Profesional I Instrumento Rúbrica 20% Portfolio 40% Entrevista 40% saberes S2 S3 P1",
        "p2": "Práctica Profesional II Instrumento Rúbrica 20% Bitácora 35% Defensa 45% saberes S1 S7 P2",
        "otra": "El estudiante debe presentar el seguro estudiantil antes de iniciar la práctica en la empresa",
    }
    r = rasgos_distintivos(textos, umbral=0.15, maximo=5)
    assert set(r) == {"p1", "p2"} and r["p1"]["grupo_similar"] == r["p2"]["grupo_similar"] == "p1"
    assert "Portfolio" in r["p1"]["rasgos_distintivos"] and "Defensa" in r["p2"]["rasgos_distintivos"]
    assert "II" in r["p2"]["rasgos_distintivos"] or "P2" in r["p2"]["rasgos_distintivos"]
    assert not set(r["p1"]["rasgos_distintivos"]) & set(r["p2"]["rasgos_distintivos"])
    assert rasgos_distintivos(textos, umbral=0.15, maximo=5) == r


def test_sin_hermanas_no_hay_grupos_ni_rasgos():
    textos = {"a": "garantía del producto y reclamos", "b": "calendario de pagos trimestrales del fondo"}
    assert agrupar_hermanos(textos, 0.5)[0] == []
    assert rasgos_distintivos(textos, 0.5, 3) == {}
    assert rasgos_distintivos({"a": "x y z", "b": "x y z"}, 0.5, 0) == {}
