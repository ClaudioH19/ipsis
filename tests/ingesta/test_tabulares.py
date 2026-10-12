"""Inventarios reales en formato de archivo: identidad, columnas y fuente literal."""
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED

from openpyxl import Workbook
from openpyxl.worksheet.table import Table
import pytest

from ipsis.ingesta.parsear import parsear_documento
from ipsis.ingesta.tabulares import preparar_tabular
from ipsis.ingesta.chunking import chunks_tabulares
from ipsis.ingesta.unidades import unidades_documento
from ipsis.indice.presentacion import bloques_fuente
from ipsis.indice.contextos import RecuperadorConContextos
from ipsis.nucleo.esquemas import RetrievalCandidate
from ipsis.servicio.api import Fragmento


def leer(ruta):
    return parsear_documento(str(ruta), "inventario", "Inventario", "1")


def crear_xlsx(tmp_path, tabla=True):
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Bodega"
    hoja.append(["SKU", "Producto", "Stock", "Precio", "Notas"])
    hoja.append([7, "Taladro", 0, 18990, "Control | seguridad\nDos velocidades"])
    hoja.append([8, "Sierra", 5, 29990, None])
    hoja["A2"].number_format = hoja["A3"].number_format = "0000"
    if tabla:
        hoja.add_table(Table(displayName="Productos", ref="A1:E3"))
    ruta = tmp_path / "inventario.xlsx"
    libro.save(ruta)
    libro.close()
    return ruta


def test_xlsx_conserva_registros_completos_vacios_ceros_y_coordenadas(tmp_path):
    doc = leer(crear_xlsx(tmp_path))
    chunks = chunks_tabulares(doc)
    assert len(chunks) == 2
    assert "0007 | Taladro | 0 | 18990" in chunks[0].texto
    assert "Sierra" not in chunks[0].texto and "Taladro" not in chunks[1].texto
    for chunk in chunks:
        assert doc.texto_completo[chunk.offset_inicio:chunk.offset_fin] == chunk.texto
        tabla = bloques_fuente(doc, chunk.offset_inicio, chunk.offset_fin)[0]
        assert tabla["kind"] == "table" and len(tabla["rows"]) == 2
        assert tabla["rows"][0]["header"] is True
        assert len(tabla["rows"][1]["cells"]) == 5
        assert tabla["rows"][1]["cells"][2]["locator"].startswith("'Bodega'!C")
    assert tabla["rows"][1]["cells"][4]["text"] == ""


def test_estructura_no_divide_fila_ni_ofrece_encabezados_sueltos(tmp_path):
    preparado = preparar_tabular(crear_xlsx(tmp_path), doc_id="inv", doc_titulo="Inventario")
    hijos, padres = unidades_documento(preparado, max_palabras=1, palabras_hijo=1)
    assert len(hijos) == len(padres) == 2
    assert all("SKU" not in h.texto_fuente for h in hijos)
    r = RecuperadorConContextos(None, padres, {preparado.doc_id: preparado})
    cs, _ = r.ampliar_contextos([RetrievalCandidate(chunk=hijos[0])], 12000)
    ampliado = r.ampliar_resultado(cs[0], 12000, vecinos=3).chunk
    assert "SKU" in ampliado.texto_fuente and "Taladro" in ampliado.texto_fuente
    assert "Sierra" not in ampliado.texto_fuente
    tabla = bloques_fuente(preparado, ampliado.offset_inicio, ampliado.offset_fin)[0]
    assert tabla["rows"][1]["cells"][0]["locator"] == "'Bodega'!A2"
    Fragmento.model_validate({"doc_title": "Inventario", "section": "Bodega",
        "text": ampliado.texto_fuente, "source_blocks": [tabla], "chunk_id": ampliado.chunk_id,
        "version": "1", "offset_start": ampliado.offset_inicio, "offset_end": ampliado.offset_fin, "score": .5})
    for padre in padres.values():
        assert preparado.texto_completo[padre.offset_inicio:padre.offset_fin] == padre.texto_fuente


@pytest.mark.parametrize("ext,contenido", [
    ("csv", 'SKU;Producto;Stock;Notas\n0007;Taladro;0;"Línea uno\nLínea dos"\n0008;Sierra;5;\n'),
    ("tsv", 'SKU\tProducto\tStock\tNotas\n0007\tTaladro\t0\t"Línea uno\nLínea dos"\n0008\tSierra\t5\t\n'),
])
@pytest.mark.parametrize("salto", ["\n", "\r\n"])
def test_csv_tsv_conservan_identificadores_ceros_y_celdas_multilinea(tmp_path, ext, contenido, salto):
    ruta = tmp_path / f"inventario.{ext}"
    ruta.write_bytes(contenido.replace("\n", salto).encode("utf-8-sig"))
    doc = leer(ruta)
    chunks = chunks_tabulares(doc)
    assert len(chunks) == 2 and "0007" in chunks[0].texto
    b = bloques_fuente(doc, chunks[0].offset_inicio, chunks[0].offset_fin)[0]
    assert b["rows"][1]["cells"][3]["text"] == f"Línea uno{salto}Línea dos"
    assert b["rows"][1]["cells"][2]["locator"] == "fila 2, columna C"


def test_csv_windows_y_filas_irregulares_conserva_columnas(tmp_path):
    ruta = tmp_path / "inventario.csv"
    ruta.write_bytes("SKU,Producto,Stock\n0007,Taladró\n0008,Sierra,5,Extra\n".encode("cp1252"))
    d = leer(ruta)
    assert any("Windows-1252" in a for a in d.advertencias)
    c = chunks_tabulares(d)[0]
    b = bloques_fuente(d, c.offset_inicio, c.offset_fin)[0]
    assert b["columns"] == 4
    assert b["rows"][1]["cells"][2]["text"] == ""
    assert "Extra" in d.texto_completo


def test_xls_binario_real():
    ruta = Path(__file__).resolve().parents[1] / "fixtures/inventario_sintetico.xls"
    d = leer(ruta)
    cs = chunks_tabulares(d)
    assert len(cs) == 2 and "0007 | Taladro | 0 | 18990" in cs[0].texto
    assert "#DIV/0!" in cs[1].texto
    assert "'Bodega'!D2" == bloques_fuente(d, cs[0].offset_inicio, cs[0].offset_fin)[0]["rows"][1]["cells"][3]["locator"]


def test_salto_literal_al_final_de_celda_no_se_recorta(tmp_path):
    ruta = tmp_path / "inventario.csv"
    ruta.write_bytes(b'SKU,Notas\n0007,"Cierre\n"\n')
    d = leer(ruta)
    ch = chunks_tabulares(d)[0]
    assert ch.texto.endswith("Cierre\n")
    tabla = bloques_fuente(d, ch.offset_inicio, ch.offset_fin)[0]
    assert tabla["kind"] == "table" and tabla["rows"][1]["cells"][1]["text"] == "Cierre\n"


@pytest.mark.parametrize("cache", [None, "37980"])
def test_formulas_se_recuperan_sin_ejecutar_y_avisan(tmp_path, cache):
    ruta = crear_xlsx(tmp_path)
    from openpyxl import load_workbook
    libro = load_workbook(ruta)
    libro.active["D2"] = "=18990*2"
    libro.save(ruta)
    libro.close()
    if cache:
        with ZipFile(ruta) as z:
            partes = {n: z.read(n) for n in z.namelist()}
        import re
        partes["xl/worksheets/sheet1.xml"] = re.sub(
            rb'(<c r="D2"[^>]*>.*?<v>)[^<]*(</v>)',
            lambda m: m[1] + cache.encode() + m[2], partes["xl/worksheets/sheet1.xml"])
        with ZipFile(ruta, "w", ZIP_DEFLATED) as z:
            for nombre, datos in partes.items():
                z.writestr(nombre, datos)
    d = leer(ruta)
    c = chunks_tabulares(d)[0]
    assert (cache or "=18990*2") in c.texto
    assert any("fórmula" in a for a in d.advertencias)
    celda = d.bloques[2].metadata["cells"][3]
    assert celda["formula"] == "=18990*2"


def test_hojas_notas_fuera_de_tabla_y_tabla_sin_encabezado(tmp_path):
    from openpyxl import load_workbook
    ruta = crear_xlsx(tmp_path)
    libro = load_workbook(ruta)
    libro.active["B6"] = "Nota fuera de tabla"
    hoja = libro.create_sheet("Archivo oculto")
    hoja.sheet_state = "hidden"
    hoja.append(["A", "Stock"])
    hoja.append(["Producto archivado", 0])
    hoja2 = libro.create_sheet("Sin encabezados")
    hoja2.append(["Primer producto", 7])
    hoja2.append(["Segundo producto", 8])
    hoja2.add_table(Table(displayName="SinHeader", ref="A1:B2", headerRowCount=0))
    libro.save(ruta)
    libro.close()
    d = leer(ruta)
    assert len(chunks_tabulares(d)) == 6
    assert all(s in d.texto_completo for s in ["Nota fuera de tabla", "Primer producto", "Producto archivado"])
    assert any("oculta" in a for a in d.advertencias)
    assert any("fuera de tablas" in a for a in d.advertencias)


def test_limite_celdas_y_archivo_sin_datos(tmp_path, monkeypatch):
    ruta = crear_xlsx(tmp_path)
    monkeypatch.setattr("ipsis.ingesta.tabulares.MAX_CELDAS", 2)
    with pytest.raises(ValueError, match="divide"):
        leer(ruta)
    ruta = tmp_path / "vacio.csv"
    ruta.write_text("SKU,Producto\n", encoding="utf-8")
    with pytest.raises(ValueError, match="filas de datos"):
        leer(ruta)


@pytest.mark.parametrize("estructural", [False, True])
def test_load_y_payloads_tabulares_no_dependen_del_chunker_textual(tmp_path, monkeypatch, estructural):
    from ipsis.cli import main
    from ipsis.ingesta.ingesta import parsear_corpus, construir_chunks_y_payloads, validar_payloads
    from ipsis.ingesta.metadata import cargar_taxonomia
    from ipsis.ingesta.unidades import construir_payloads_estructurados

    ruta = crear_xlsx(tmp_path)
    proyecto = tmp_path / "proyecto"
    proyecto.mkdir()
    monkeypatch.chdir(proyecto)
    assert main(["init", "--profile", "light"]) == 0
    assert main(["load", str(ruta)]) == 0
    docs = parsear_corpus("documents", "estructural_v1" if estructural else "legacy")
    taxonomia = cargar_taxonomia()
    if estructural:
        cs, payloads, _, _ = construir_payloads_estructurados(docs, taxonomia,
            {"modo": "estructural_v1", "max_palabras_unidad": 1, "palabras_hijo": 1})
    else:
        cs, payloads = construir_chunks_y_payloads(docs, "CH-F512", True, taxonomia)
        assert all(p["estrategia_chunking"] == "CH-ROW1" for p in payloads)
    assert len(cs) == len(payloads) == 2
    assert not validar_payloads(payloads, taxonomia)
    assert all(not ("Taladro" in p["texto"] and "Sierra" in p["texto"]) for p in payloads)


@pytest.mark.slow
@pytest.mark.modelos("intfloat/multilingual-e5-small")
@pytest.mark.parametrize("estructural", [False, True])
def test_inventario_hasta_fuentes_v2_con_indice_y_modelo_reales(tmp_path, monkeypatch, estructural):
    from ipsis.cli import main
    from ipsis.aplicacion.consulta import construir_motor_del_proyecto
    from ipsis.aplicacion.preparacion import preparar_proyecto
    from ipsis.servicio.api import respuesta_publica, RespuestaConsulta

    ruta = crear_xlsx(tmp_path)
    proyecto = tmp_path / "original"
    proyecto.mkdir()
    monkeypatch.chdir(proyecto)
    assert main(["init", "--profile", "light"]) == 0
    assert main(["load", str(ruta)]) == 0
    if estructural:
        copia = tmp_path / "estructural"
        preparar_proyecto(proyecto, copia)
        monkeypatch.chdir(copia)
    assert main(["index"]) == 0
    m = construir_motor_del_proyecto()
    try:
        for pregunta in ("stock del Taladro", "producto SKU 0007"):
            publica = respuesta_publica(m.consultar(pregunta, pregunta).contrato_http)
            RespuestaConsulta.model_validate(publica)
            assert publica["response_mode"] == "ranked_sources"
            f = publica["fragments"][0]
            assert "0007" in f["text"] and "Taladro" in f["text"]
            assert "Sierra" not in f["text"]
            tabla = next(b for b in f["source_blocks"] if b["kind"] == "table")
            assert tabla["rows"][1]["cells"][2]["text"] == "0"
            assert tabla["rows"][1]["cells"][2]["locator"] == "'Bodega'!C2"
    finally:
        m.cerrar()
