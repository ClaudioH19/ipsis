"""Fidelidad estructural y literal: celdas vacías, merges, orden y procedencia."""
import json

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import pytest

from ipsis.ingesta.preparacion import preparar_docx
from ipsis.ingesta.unidades import unidades_documento


def preparar(doc, tmp_path, nombre="source.docx"):
    archivo = tmp_path / nombre
    doc.save(archivo)
    return preparar_docx(archivo, doc_id="manual", doc_titulo="Manual", version="1")


def ejemplo():
    d = Document()
    d.add_heading("Evaluación", 1)
    for titulo, valores in (("Práctica I", ["20%", "40%", "40%"]), ("Práctica II", ["20%", "35%", "45%"] )):
        d.add_paragraph(titulo)
        tabla = d.add_table(rows=4, cols=3)
        tabla.cell(0, 0).text, tabla.cell(0, 1).text = "Instrumento", "Peso"
        for i, v in enumerate(valores, 1):
            tabla.cell(i, 0).text, tabla.cell(i, 1).text = f"Instrumento {i}", v
    return d


def test_tablas_vacias_y_padres_no_mezclan_sujetos(tmp_path):
    d = preparar(ejemplo(), tmp_path)
    assert len(d.tablas) == 2
    assert d.tablas[0].columnas == 3
    assert len(d.tablas[0].celdas) == 12
    assert d.tablas[0].celdas[2].texto == ""
    assert "Instrumento | Peso | " in d.texto_completo
    cs, padres = unidades_documento(d)
    p1 = next(p for p in padres.values() if p.texto_fuente.startswith("Práctica I\n"))
    assert "20%" in p1.texto_fuente and p1.texto_fuente.count("40%") == 2
    assert "35%" not in p1.texto_fuente and "Práctica II" not in p1.texto_fuente
    for ch in [*cs, *padres.values()]:
        assert d.texto_completo[ch.offset_inicio:ch.offset_fin] == ch.texto_fuente
    for hijo in cs:
        padre = padres[hijo.parent_id]
        assert padre.offset_inicio <= hijo.offset_inicio <= hijo.offset_fin <= padre.offset_fin
    assert all(not t.encabezado_explicito for t in d.tablas)


def test_combinaciones_horizontales_y_verticales_no_replican_texto(tmp_path):
    doc = Document()
    t = doc.add_table(rows=3, cols=3)
    t.cell(0, 0).merge(t.cell(1, 0)).text = "Origen vertical"
    t.cell(2, 0).merge(t.cell(2, 1)).text = "Origen horizontal"
    t.cell(2, 2).text = "Final"
    d = preparar(doc, tmp_path)
    tabla = d.tablas[0]
    vertical = next(c for c in tabla.celdas if c.texto == "Origen vertical")
    horizontal = next(c for c in tabla.celdas if c.texto == "Origen horizontal")
    assert vertical.filas == 2 and horizontal.columnas == 2
    continuacion = next(c for c in tabla.celdas if c.continuacion_de == vertical.cell_id)
    assert continuacion.texto == "" and continuacion.fila == 1
    assert d.texto_completo.count("Origen vertical") == 1
    for celda in tabla.celdas:
        r = celda.referencia
        assert d.texto_completo[r.offset_inicio:r.offset_fin] == celda.texto


def test_celdas_omitidas_no_desplazan_columnas(tmp_path):
    doc = Document()
    t = doc.add_table(rows=1, cols=3)
    t.cell(0, 1).text = "Columna B"
    t.cell(0, 2).text = "Columna C"
    fila = t.rows[0]._tr
    fila.remove(fila.tc_lst[0])
    antes = OxmlElement("w:gridBefore")
    antes.set(qn("w:val"), "1")
    fila.get_or_add_trPr().append(antes)
    d = preparar(doc, tmp_path)
    assert [c.columna for c in d.tablas[0].celdas] == [1, 2]
    assert d.texto_completo.startswith(" | Columna B | Columna C")
    fila_elemento = next(e for e in d.elementos if e.metadata.get("tipo") == "fila")
    assert fila_elemento.metadata["omitidas_inicio"] == 1


def test_huella_estructural_cambia_aunque_el_texto_sea_igual(tmp_path):
    a, b = Document(), Document()
    ta, tb = a.add_table(rows=1, cols=2), b.add_table(rows=1, cols=2)
    ta.cell(0, 0).merge(ta.cell(0, 1)).text = "A"
    tb.cell(0, 0).text = "A"
    da, db = preparar(a, tmp_path, "a.docx"), preparar(b, tmp_path, "b.docx")
    assert da.texto_completo == db.texto_completo
    assert da.structure_hash != db.structure_hash
    assert preparar_docx(tmp_path / "a.docx", doc_id="manual", doc_titulo="Manual").structure_hash == da.structure_hash


def test_contenido_anidado_y_avisos_no_se_ocultan(tmp_path):
    doc = Document()
    celda = doc.add_table(rows=1, cols=1).cell(0, 0)
    celda.text = "Antes"
    celda.add_table(rows=1, cols=2).cell(0, 0).text = "Contenido interno"
    celda.add_paragraph("Después")
    d = preparar(doc, tmp_path)
    assert d.texto_completo.index("Antes") < d.texto_completo.index("Contenido interno") < d.texto_completo.index("Después")
    assert any("anidada" in aviso for aviso in d.advertencias)


def test_ventanas_hijo_cubren_parrafo_sin_perdida(tmp_path):
    doc = Document()
    texto = " ".join(f"palabra{i}" for i in range(200))
    doc.add_paragraph(texto)
    d = preparar(doc, tmp_path)
    cs, padres = unidades_documento(d, palabras_hijo=30)
    assert "".join(c.texto_fuente for c in cs) == texto
    assert len(padres) == 1
    assert all(len(c.texto_fuente.split()) <= 30 for c in cs)


def test_formato_marca_limite_sin_inferir_intencion(tmp_path):
    doc = Document()
    doc.add_paragraph("Texto anterior.")
    doc.add_paragraph().add_run("Grupo A").bold = True
    doc.add_paragraph("Contenido A.")
    doc.add_paragraph().add_run("Grupo B").bold = True
    doc.add_paragraph("Contenido B.")
    d = preparar(doc, tmp_path)
    _, padres = unidades_documento(d)
    assert any(p.texto_fuente == "Grupo A\nContenido A." for p in padres.values())
    assert not any("Grupo A" in p.texto_fuente and "Grupo B" in p.texto_fuente for p in padres.values())
    assert all(not e.seccion_path for e in d.elementos)  # negrita no inventa nivel de sección


def test_tabla_grande_se_divide_y_conserva_referencias_a_encabezados(tmp_path):
    doc = Document()
    t = doc.add_table(rows=30, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Pieza", "Valor"
    for i in range(1, 30):
        t.cell(i, 0).text, t.cell(i, 1).text = f"Pieza {i}", str(i)
    d = preparar(doc, tmp_path)
    cs, ps = unidades_documento(d, max_palabras=20)
    assert len(ps) > 1
    assert len(cs) == 30
    assert all(p.metadata["referencias_contexto"] for p in ps.values())
    assert all(d.texto_completo[p.offset_inicio:p.offset_fin] == p.texto_fuente for p in ps.values())


def test_campos_de_unidad_posterior_conservan_cabecera_y_referencias(tmp_path):
    from ipsis.nucleo.lectura import ContextoLectura

    doc = Document()
    doc.add_paragraph("Inventario original")
    tabla = doc.add_table(rows=8, cols=2)
    for i in range(8):
        tabla.cell(i, 0).text = "Pieza" if i == 0 else f"Pieza {i}"
        tabla.cell(i, 1).text = "Valor" if i == 0 else str(i * 10)
    d = preparar(doc, tmp_path)
    _, ps = unidades_documento(d, max_palabras=8)
    ultimo = list(ps.values())[-1]
    contexto = ContextoLectura.model_validate(ultimo.metadata["contexto_lectura"])
    assert "Pieza: Pieza 7; Valor: 70" in contexto.texto
    assert "Inventario original" in contexto.texto
    assert "Inventario original" not in ultimo.texto_fuente
    contexto.verificar_fuente(d.texto_completo)
    inicio = contexto.texto.index("Valor: 70")
    refs = contexto.referencias_span(inicio, inicio + len("Valor: 70"))
    assert refs[0]["offset_fin"] <= ultimo.offset_inicio  # cabecera fuera de la unidad
    assert refs[1]["texto"] == "70" and refs[1]["cell_id"]
    corto = contexto.recortar(inicio + len("Valor: 7"))
    corto.verificar_fuente(d.texto_completo)
    assert corto.referencias_span(inicio, len(corto.texto))[-1]["texto"] == "7"


def test_encabezado_word_declarado_no_se_presenta_como_inferido(tmp_path):
    from ipsis.nucleo.lectura import ContextoLectura

    doc = ejemplo()
    for t in doc.tables:
        t.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    d = preparar(doc, tmp_path)
    _, ps = unidades_documento(d)
    for padre in ps.values():
        if padre.tipo_contenido == "tabla":
            contexto = ContextoLectura.model_validate(padre.metadata["contexto_lectura"])
            assert not contexto.supuestos
            assert any(c.rol == "encabezado" for c in contexto.copias)
