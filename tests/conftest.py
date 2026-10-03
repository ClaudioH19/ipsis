"""Fixtures globales basadas únicamente en datos sintéticos."""
from __future__ import annotations

import os
from pathlib import Path

import pytest
from docx import Document

from rag.lenguaje.dominio import cargar_perfil
from rag.lenguaje.entidades import configurar_idioma, configurar_perfil
from rag.lenguaje.paquete import cargar_paquete


@pytest.hookimpl(tryfirst=True)
def pytest_configure(config: pytest.Config) -> None:
    """Elige un directorio temporal que pytest pueda crear y borrar.

    En Windows, el temporal del sistema puede devolver "Acceso denegado" y una
    carpeta dentro del checkout falla si está sincronizada por OneDrive, que
    bloquea el borrado entre corridas. Se usa ``%LOCALAPPDATA%``, fuera de
    ambos. ``RAG_PYTEST_BASETEMP`` o ``--basetemp`` tienen prioridad; en otros
    sistemas se conserva el comportamiento por defecto de pytest.
    """
    if config.option.basetemp:
        return
    elegido = os.environ.get("RAG_PYTEST_BASETEMP")
    if not elegido and os.name == "nt" and os.environ.get("LOCALAPPDATA"):
        elegido = str(Path(os.environ["LOCALAPPDATA"]) / "rag-extractivo" / "pytest")
    if elegido:
        config.option.basetemp = elegido


def ruta_cache_modelos() -> Path:
    """Caché global de modelos; debe coincidir con ``rag.modelos.descarga``."""
    elegido = os.environ.get("RAG_MODELOS_DIR")
    if elegido:
        return Path(elegido)
    return Path.home() / ".cache" / "rag-extractivo" / "modelos"


def _modelo_en_cache(id_hf: str) -> bool:
    # Estructura de caché de huggingface_hub: models--<org>--<nombre>/snapshots/
    carpeta = ruta_cache_modelos() / ("models--" + id_hf.replace("/", "--")) / "snapshots"
    return carpeta.is_dir() and any(carpeta.iterdir())


def pytest_collection_modifyitems(config: pytest.Config, items: list[pytest.Item]) -> None:
    """Omite las pruebas ``modelos`` cuyos modelos no están descargados."""
    for item in items:
        marca = item.get_closest_marker("modelos")
        if marca is None:
            continue
        faltantes = [id_hf for id_hf in marca.args if not _modelo_en_cache(id_hf)]
        if faltantes or not marca.args:
            item.add_marker(pytest.mark.skip(
                reason=f"faltan modelos en caché: {faltantes or 'sin id'}; ejecuta 'rag modelos descargar'"
            ))


@pytest.fixture(scope="session")
def proyecto_suite(tmp_path_factory: pytest.TempPathFactory) -> Path:
    """Proyecto creado con ``rag init``, compartido por la suite.

    Reemplaza a la antigua configuración en la raíz del repositorio: las
    pruebas leen ``configuracion/`` y ``documentos/`` relativos al directorio
    actual, igual que el servicio real.
    """
    from rag.cli import main

    carpeta = tmp_path_factory.mktemp("proyecto_suite")
    anterior = Path.cwd()
    os.chdir(carpeta)
    try:
        assert main(["init"]) == 0
    finally:
        os.chdir(anterior)
    return carpeta


@pytest.fixture(autouse=True)
def _en_proyecto_suite(proyecto_suite: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.chdir(proyecto_suite)


@pytest.fixture(scope="session", autouse=True)
def _perfil_generico_de_la_suite(proyecto_suite: Path):
    configurar_perfil(cargar_perfil(str(proyecto_suite / "configuracion" / "dominio.yaml")))
    configurar_idioma(cargar_paquete("es"))
    yield
    configurar_perfil(None)
    configurar_idioma(None)


@pytest.fixture(scope="session")
def ruta_docx_sintetico(tmp_path_factory: pytest.TempPathFactory) -> Path:
    ruta = tmp_path_factory.mktemp("documentos_sinteticos") / "manual_x100.docx"
    doc = Document()
    doc.add_heading("Manual del Equipo X100", level=1)
    doc.add_paragraph("Documento ficticio creado exclusivamente para las pruebas.")
    doc.add_heading("Operación", level=1)
    doc.add_paragraph("Para encender el equipo X100, mantén presionado el botón verde durante 3 segundos.")
    doc.add_heading("Mantenimiento", level=1)
    doc.add_paragraph("Limpia el filtro cada 30 días, salvo que el entorno tenga mucho polvo.")
    tabla = doc.add_table(rows=2, cols=2)
    tabla.cell(0, 0).text = "Componente"
    tabla.cell(0, 1).text = "Intervalo"
    tabla.cell(1, 0).text = "Filtro"
    tabla.cell(1, 1).text = "30 días"
    doc.save(ruta)
    return ruta
