"""La vista conserva celdas, relaciones declaradas y offsets del texto fuente."""
from docx import Document
from docx.oxml import OxmlElement

from ipsis.indice.presentacion import bloques_fuente
from ipsis.indice.contextos import RecuperadorConContextos
from ipsis.ingesta.parsear import parsear_docx, parsear_markdown
from ipsis.ingesta.preparacion import preparar_docx
from ipsis.servicio.api import RespuestaConsulta, respuesta_publica
from tests.consulta.test_motor_v2 import candidato, motor, Recuperador


def documento_tabla(tmp_path):
    doc = Document()
    doc.add_heading("Evaluación", 1)
    doc.add_paragraph("Primera explicación.\nSegunda línea.")
    tabla = doc.add_table(rows=3, cols=3)
    tabla.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for celda, texto in zip(tabla.rows[0].cells, ["Instrumento", "Peso", "Observaciones"]):
        celda.text = texto
    tabla.cell(1, 0).text = "Control | seguridad 😊 <img src=x onerror=alert(1)>"
    tabla.cell(1, 1).text = "20%"
    tabla.cell(2, 0).text = "Informe"
    tabla.cell(2, 1).text = "80%"
    tabla.cell(2, 2).text = "Revisión final"
    archivo = tmp_path / "manual.docx"
    doc.save(archivo)
    return archivo


def test_docx_legacy_no_cambia_texto_y_recupera_columnas_vacias(tmp_path):
    d = parsear_docx(str(documento_tabla(tmp_path)), "reglamento", "Manual", "1")
    b = bloques_fuente(d, 0, len(d.texto_completo))
    tabla = next(x for x in b if x["kind"] == "table")
    assert tabla["columns"] == 3 and len(tabla["rows"]) == 3
    assert tabla["rows"][0]["header"] is True
    assert tabla["rows"][1]["cells"][0]["text"].startswith("Control | seguridad")
    assert tabla["rows"][1]["cells"][2]["text"] == ""
    assert "20%\nInforme | 80% | Revisión final" in d.texto_completo
    for fila in tabla["rows"]:
        for c in fila["cells"]:
            assert d.texto_completo[c["offset_start"]:c["offset_end"]] == c["text"]


def test_tabla_preparada_conserva_combinaciones_y_encabezados_no_inferidos(tmp_path):
    doc = Document()
    t = doc.add_table(rows=3, cols=3)
    t.cell(0, 0).merge(t.cell(1, 0)).text = "Vertical"
    t.cell(2, 0).merge(t.cell(2, 1)).text = "Horizontal"
    archivo = tmp_path / "combinadas.docx"
    doc.save(archivo)
    d = preparar_docx(archivo, doc_id="d", doc_titulo="Combinadas")
    tabla = bloques_fuente(d, 0, len(d.texto_completo))[0]
    assert tabla["rows"][0]["cells"][0]["row_span"] == 2
    assert tabla["rows"][2]["cells"][0]["column_span"] == 2
    assert not any(f["header"] for f in tabla["rows"])
    recortado = bloques_fuente(d, 1, len(d.texto_completo) - 1)
    assert all(b["kind"] != "table" for b in recortado)
    assert recortado[0]["text"] == d.texto_completo[1:-1]


def test_markdown_mantiene_posiciones_y_reconoce_encabezado_declarado(tmp_path):
    archivo = tmp_path / "inventario.md"
    archivo.write_text("# Inventario\n\n| Producto | Stock | Color |\n| --- | --- | --- |\n"
                       "| A | | Azul |\n| B | 5 | |\n", encoding="utf-8")
    d = parsear_markdown(str(archivo), "d", "Inventario", "1")
    tabla = next(b for b in bloques_fuente(d, 0, len(d.texto_completo)) if b["kind"] == "table")
    assert tabla["columns"] == 3 and tabla["rows"][0]["header"]
    assert tabla["rows"][1]["cells"][1]["text"] == ""
    assert tabla["rows"][1]["cells"][2]["text"] == "Azul"


def test_encabezado_docx_desactivado_no_se_infiere(tmp_path):
    from docx.oxml.ns import qn

    archivo = documento_tabla(tmp_path)
    original = Document(archivo)
    original.tables[0].rows[0]._tr.trPr.find(qn("w:tblHeader")).set(qn("w:val"), "0")
    original.save(archivo)
    d = parsear_docx(str(archivo), "d", "Manual", "1")
    tabla = next(b for b in bloques_fuente(d, 0, len(d.texto_completo)) if b["kind"] == "table")
    assert not tabla["rows"][0]["header"]


def test_api_v2_publica_tabla_con_primera_fila_fuera_de_ruta_legacy(tmp_path):
    d = parsear_docx(str(documento_tabla(tmp_path)), "reglamento", "Manual", "1")
    fila = next(b for b in d.bloques if b.tipo_contenido == "tabla" and "20%" in b.texto)
    c = candidato(fila.texto, inicio=fila.offset_inicio, seccion=fila.seccion_path)
    m = motor(candidatos=[c])
    m.retriever = RecuperadorConContextos(Recuperador([c]), {}, {d.doc_id: d})
    publica = respuesta_publica(m.consultar("peso seguridad", "s").contrato_http)
    RespuestaConsulta.model_validate(publica)
    f = publica["fragments"][0]
    assert f["text"] == d.texto_completo[f["offset_start"]:f["offset_end"]]
    tabla = next(b for b in f["source_blocks"] if b["kind"] == "table")
    assert tabla["rows"][0]["cells"][0]["text"] == "Instrumento"
    assert f["match_start"] == fila.offset_inicio
