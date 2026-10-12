"""La expansión no inventa texto ni mezcla versiones; almacenes dañados fallan."""
from types import SimpleNamespace

import pytest

from ipsis.indice.contextos import RecuperadorConContextos, cargar_contextos, guardar_contextos, ruta_contextos
from ipsis.indice.manifiesto import ManifiestoIndice, hash_corpus
from ipsis.indice.payload import chunk_desde_payload, payload_desde_chunk
from ipsis.ingesta.preparacion import preparar_docx
from ipsis.ingesta.unidades import unidades_documento
from ipsis.nucleo.esquemas import RetrievalCandidate, RetrievalSignals
from docx import Document


def unidades(tmp_path):
    doc = Document()
    doc.add_paragraph("Contexto fuente.")
    tabla = doc.add_table(rows=3, cols=2)
    tabla.cell(0, 0).text, tabla.cell(0, 1).text = "Concepto", "Peso"
    tabla.cell(1, 0).text, tabla.cell(1, 1).text = "Parte A", "20%"
    tabla.cell(2, 0).text, tabla.cell(2, 1).text = "Parte B", "80%"
    archivo = tmp_path / "fuente.docx"
    doc.save(archivo)
    d = preparar_docx(archivo, doc_id="fuente", doc_titulo="Fuente")
    cs, ps = unidades_documento(d)
    return d, cs, ps


def candidato(ch, rango=1):
    return RetrievalCandidate(chunk=ch, signals=RetrievalSignals(fusion_rank=rango, fusion_score=1 / (60 + rango)))


def test_expande_dos_filas_a_una_tabla_sin_sumar_votos(tmp_path):
    _, cs, ps = unidades(tmp_path)
    filas = [c for c in cs if c.tipo_contenido == "tabla"][1:]
    r = RecuperadorConContextos(SimpleNamespace(), ps)
    salida, eventos = r.ampliar_contextos([candidato(filas[0]), candidato(filas[1], 2)], 12000)
    assert len(salida) == 1
    assert "20%" in salida[0].chunk.texto_fuente and "80%" in salida[0].chunk.texto_fuente
    assert salida[0].score == candidato(filas[0]).score
    assert len(salida[0].chunk.metadata["hijos_recuperados"]) == 2
    assert all(e["aplicada"] for e in eventos)
    assert "hijos_recuperados" not in next(iter(ps.values())).metadata


@pytest.mark.parametrize("campo,valor", [("version", "otra"), ("texto_fuente", "texto ajeno")])
def test_rechaza_hijo_de_otra_fuente(tmp_path, campo, valor):
    _, cs, ps = unidades(tmp_path)
    corrupto = cs[-1].model_copy(update={campo: valor})
    with pytest.raises(ValueError, match="otra extracción"):
        RecuperadorConContextos(None, ps).ampliar_contextos([candidato(corrupto)], 12000)


def test_no_pierde_hijo_si_ampliar_lo_dejaria_fuera_del_presupuesto(tmp_path):
    _, cs, ps = unidades(tmp_path)
    salida, eventos = RecuperadorConContextos(None, ps).ampliar_contextos([candidato(cs[-1])], 8)
    assert salida[0].chunk == cs[-1]
    assert eventos[0]["aplicada"] is False


def test_contexto_de_titulo_incluye_tabla_completa_literal(tmp_path):
    doc, _, padres = unidades(tmp_path)
    parrafo = next(p for p in padres.values() if p.tipo_contenido != "tabla")
    r = RecuperadorConContextos(None, padres, {doc.doc_id: doc})
    c = r.ampliar_resultado(candidato(parrafo), 16000, vecinos=1)
    assert "20%" in c.chunk.texto_fuente and "80%" in c.chunk.texto_fuente
    assert doc.texto_completo[c.chunk.offset_inicio:c.chunk.offset_fin] == c.chunk.texto_fuente
    assert c.chunk.metadata["ventana_resultado"]["coincidencia_inicio"] == parrafo.offset_inicio
    assert parrafo.texto_fuente == "Contexto fuente."


def test_contexto_se_detiene_en_limite_de_seccion_y_version():
    from tests.consulta.test_motor_v2 import candidato as crear

    a, b, c = [crear(t, str(i), inicio=2*i).chunk for i, t in enumerate(["A", "B", "C"])]
    b.seccion_path = ["Otra"]
    r = RecuperadorConContextos(None, {ch.chunk_id: ch for ch in (a, b, c)},
        {a.doc_id: SimpleNamespace(version="1", texto_completo="A\nB\nC")})
    assert r.ampliar_resultado(candidato(a), 256, vecinos=2).chunk.texto_fuente == "A"
    b.seccion_path, b.version = a.seccion_path, "2"
    assert r.ampliar_resultado(candidato(a), 256, vecinos=2).chunk.texto_fuente == "A"


def test_contexto_legacy_centra_una_coincidencia_tardia_y_no_cruza_extracciones():
    from tests.consulta.test_motor_v2 import candidato as crear

    fuente = "Antes. " * 200 + "Producto X: stock 7." + " Después." * 200
    ch = crear("Producto X: stock 7.", inicio=fuente.index("Producto")).chunk
    bloque = SimpleNamespace(offset_inicio=0, offset_fin=len(fuente), seccion_path=ch.seccion_path)
    doc = SimpleNamespace(version="1", texto_completo=fuente, bloques=[bloque])
    r = RecuperadorConContextos(None, {}, {ch.doc_id: doc})
    resultado = r.ampliar_resultado(candidato(ch), 256)
    assert len(resultado.chunk.texto_fuente) <= 256 and ch.texto_fuente in resultado.chunk.texto_fuente
    assert fuente[resultado.chunk.offset_inicio:resultado.chunk.offset_fin] == resultado.chunk.texto_fuente
    assert resultado.chunk.metadata["ventana_resultado"]["truncado"]
    doc.texto_completo = fuente.replace("stock 7", "stock 8")
    with pytest.raises(ValueError, match="extracción"):
        r.ampliar_resultado(candidato(ch), 256)


def test_almacen_versionado_se_verifica_con_manifiesto(tmp_path):
    d, _, ps = unidades(tmp_path)
    huella = guardar_contextos("nueva", [d], ps, {"modo": "estructural_v1"}, tmp_path)
    m = ManifiestoIndice(coleccion="nueva", modelo_denso="x", dimension=2, hash_corpus="x", n_chunks=3,
                        preparacion={"contextos_sha256": huella})
    assert cargar_contextos(m, tmp_path) == ps
    ruta_contextos("nueva", tmp_path).write_text("{}")
    with pytest.raises(ValueError, match="manifiesto"):
        cargar_contextos(m, tmp_path)


def test_hash_indice_cubre_relaciones_y_no_cambia_legacy():
    a = {"hash_texto": "x", "preparacion": {"modo": "estructural_v1"}, "parent_id": "p1"}
    assert hash_corpus([a]) != hash_corpus([{**a, "parent_id": "p2"}])
    assert hash_corpus([{"hash_texto": "x"}]) == hash_corpus([{"hash_texto": "x", "parent_id": "p2"}])


def test_adaptador_conserva_padre_y_vecinos(tmp_path):
    _, cs, _ = unidades(tmp_path)
    c = cs[-1]
    payload = payload_desde_chunk(c)
    restaurado = chunk_desde_payload(payload)
    assert restaurado.parent_id == c.parent_id
    assert restaurado.prev_id == c.prev_id
    assert payload_desde_chunk(restaurado) == payload


def test_almacen_rechaza_mapa_no_literal_aun_con_checksum_valido(tmp_path):
    d, _, ps = unidades(tmp_path)
    padre = next(p for p in ps.values() if p.tipo_contenido == "tabla")
    mapa = padre.metadata["contexto_lectura"]
    copia = next(c for c in mapa["copias"] if c["cell_id"])
    inicio, fin = copia["inicio"], copia["fin"]
    mapa["texto"] = mapa["texto"][:inicio] + "X" * (fin - inicio) + mapa["texto"][fin:]
    huella = guardar_contextos("corrupta", [d], ps, {"modo": "estructural_v1"}, tmp_path)
    m = ManifiestoIndice(coleccion="corrupta", modelo_denso="x", dimension=2, hash_corpus="x", n_chunks=3,
                        preparacion={"contextos_sha256": huella})
    with pytest.raises(ValueError, match="lectura no literal"):
        cargar_contextos(m, tmp_path)


def test_almacen_rechaza_celda_de_otra_posicion(tmp_path):
    d, _, ps = unidades(tmp_path)
    padre = next(p for p in ps.values() if p.tipo_contenido == "tabla")
    copias = padre.metadata["contexto_lectura"]["copias"]
    celdas = [c for c in copias if c["cell_id"] and c["rol"] == "valor"]
    celdas[0]["cell_id"] = celdas[-1]["cell_id"]
    huella = guardar_contextos("identidad", [d], ps, {"modo": "estructural_v1"}, tmp_path)
    m = ManifiestoIndice(coleccion="identidad", modelo_denso="x", dimension=2, hash_corpus="x", n_chunks=3,
                        preparacion={"contextos_sha256": huella})
    with pytest.raises(ValueError, match="identidad de celda"):
        cargar_contextos(m, tmp_path)
