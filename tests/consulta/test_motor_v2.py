"""Invariantes del recorrido experimental y su frontera con servicio/configuración."""
from __future__ import annotations

import ast
from pathlib import Path
from types import SimpleNamespace

import pytest

from ipsis.aplicacion.registro_base import LiteralOffsetsVerifier
from ipsis.configuracion.esquemas import migrar_motor, validar_motor
from ipsis.consulta.admision import ConfiguracionAdmision, ControlAdmision
from ipsis.consulta.errores import MotorSaturado
from ipsis.consulta.v2 import MotorV2
from ipsis.nucleo.contratos import QueryEngine
from ipsis.nucleo.esquemas import AssistantState, ChunkRecord, LecturaQA, RetrievalCandidate, RetrievalSignals
from ipsis.nucleo.traza import capturar_traza


def candidato(texto, cid="table", *, seccion=None, inicio=100):
    return RetrievalCandidate(chunk=ChunkRecord(chunk_id=cid, doc_id="reglamento", doc_titulo="Reglamento",
        version="1", texto_fuente=texto, texto_embedding=texto, offset_inicio=inicio,
        offset_fin=inicio + len(texto), seccion_path=seccion or ["P1"], tipo_contenido="tabla"),
        signals=RetrievalSignals(fusion_rank=1, fusion_score=.04))


class Recuperador:
    def __init__(self, candidatos):
        self.candidatos, self.consultas = candidatos, []

    def recuperar(self, consulta, limite, filtro=None):
        assert filtro is None
        self.consultas.append((consulta, limite))
        return [c.model_copy(deep=True) for c in self.candidatos[:limite]]

    def calentar(self):
        pass

    def n_chunks(self):
        return len(self.candidatos)

    def inventario_entidades(self):
        return set()

    def estadisticas_lexico(self):
        return SimpleNamespace()


class Lector:
    def __init__(self, tramo=None, corromper=False):
        self.tramo, self.corromper, self.entradas = tramo, corromper, []

    def leer(self, pregunta, pasajes):
        self.entradas.append((pregunta, pasajes))
        salida = []
        for i, texto in enumerate(pasajes):
            ini = texto.find(self.tramo) if self.tramo else -1
            if ini >= 0:
                salida.append(LecturaQA(indice_pasaje=i, inicio=ini, fin=ini + len(self.tramo),
                    texto="XX%" if self.corromper else self.tramo, prob_respuesta=.95, prob_sin_respuesta=.05))
            else:
                salida.append(LecturaQA(indice_pasaje=i, inicio=0, fin=0, texto="",
                                       prob_respuesta=0, prob_sin_respuesta=1))
        return salida


TABLA = "Práctica Profesional I\nInstrumento | Ponderación\nSupervisor | 20%\nPortafolio | 40%\nEntrevista | 40%"


def motor(candidatos=None, lector=None, **kwargs):
    return MotorV2(retriever=Recuperador([candidato(TABLA)] if candidatos is None else candidatos),
        verificador=LiteralOffsetsVerifier(), lector=lector,
        plantillas={"abstencion": {"generico": ["No hay evidencia disponible."],
                                  "zona_gris_intro": ["Pasajes recuperados para revisar:"]}}, **kwargs)


def test_porcentaje_no_reemplaza_su_tabla_y_offsets_globales():
    lector = Lector("20%")
    m = motor(lector=lector)
    resultado = m.consultar("¿Qué ponderación tiene la evaluación?", "s")
    assert isinstance(m, QueryEngine)
    assert resultado.estado == AssistantState.RELACIONADO
    assert resultado.evidencias[0].texto == TABLA
    assert "Portafolio | 40%" in resultado.mensaje_usuario
    assert "Entrevista | 40%" in resultado.mensaje_usuario
    assert lector.entradas[0][1] == [TABLA]
    bloque = next(b for b in resultado.contrato_http["bloques"] if b["tipo"] == "pasaje")
    destacado = bloque["destacados"][0]["referencias"][0]
    assert destacado["offset_inicio"] == 100 + TABLA.index("20%")
    assert "respuesta_exacta" not in resultado.contrato_http
    assert "nivel_confianza" not in resultado.contrato_http


def test_sin_span_muestra_tabla_literal_como_recuperada():
    resultado = motor(lector=Lector()).consultar("¿Cuáles son las ponderaciones?", "s")
    assert resultado.estado == AssistantState.RELACIONADO
    assert resultado.evidencias[0].texto == TABLA
    assert "respuesta_exacta" not in resultado.contrato_http
    assert resultado.diagnostico["regla"] == "recuperado"


def test_presupuesto_contexto_es_visible_y_no_falsea_offsets():
    fuente = TABLA + "\n" + "Contenido adicional. " * 30
    resultado = motor(candidatos=[candidato(fuente)], opciones={"max_caracteres_contexto": 256,
        "max_caracteres_resultado": 256}).consultar("tabla", "s")
    evidencia = resultado.evidencias[0]
    assert evidencia.texto == fuente[:256]
    assert evidencia.offset_fin == evidencia.offset_inicio + 256
    assert resultado.diagnostico["senales"]["contextos"][0]["truncado"] is True
    assert resultado.diagnostico["advertencias"]


def test_diagnostico_v2_no_reinterpreta_ni_repite_busquedas():
    from ipsis.aplicacion.consulta import diagnosticar_pregunta

    m = motor()

    def prohibido(*_):
        raise AssertionError("el diagnóstico v2 debe usar la consulta observada")

    m.retriever.rankings_diagnostico = prohibido
    m.retriever.textos_diagnostico = prohibido
    resultado = diagnosticar_pregunta("ponderaciones", m)
    assert resultado.oraciones == []
    assert len(m.retriever.consultas) == 1
    assert resultado.tramo is None
    assert resultado.regla == "recuperado"


def test_sin_candidatos_es_sin_evidencia():
    resultado = motor(candidatos=[]).consultar("Pregunta ajena", "s")
    assert resultado.estado == AssistantState.SIN_RESPUESTA
    assert not resultado.evidencias


def test_no_hereda_la_relacion_del_turno_anterior():
    lector = Lector()
    m = motor(lector=lector)
    m.consultar("¿Quién evalúa la práctica?", "same-session")
    nueva = "¿Cuándo comienza la práctica?"
    m.consultar(nueva, "same-session")
    assert lector.entradas[-1][0] == nueva
    assert m.retriever.consultas[-1][0].pregunta_original == nueva


def test_span_ajeno_falla_y_libera_admision_con_traza_parcial():
    control = ControlAdmision(ConfiguracionAdmision(concurrencia_maxima=1))
    m = motor(lector=Lector("20%", corromper=True), control_admision=control)
    with capturar_traza() as traza:
        with pytest.raises(ValueError, match="offsets"):
            m.consultar("porcentaje", "s")
    assert any(e.get("nombre") == "lectura" and e["estado"] == "error" for e in traza.eventos)
    assert control.intentar_admitir()
    control.liberar()


def test_saturacion_no_ejecuta_recuperacion():
    control = ControlAdmision(ConfiguracionAdmision(concurrencia_maxima=1))
    m = motor(control_admision=control)
    assert control.intentar_admitir()
    try:
        with pytest.raises(MotorSaturado):
            m.consultar("pregunta", "s")
        assert not m.retriever.consultas
    finally:
        control.liberar()


def test_rechaza_evidencia_corrupta_sin_publicarla():
    class Rechazar:
        def verificar(self, *_):
            return False, ["procedencia inválida"]

    m = motor(lector=Lector("20%"))
    m.verificador = Rechazar()
    resultado = m.consultar("porcentaje", "s")
    assert resultado.estado == AssistantState.SIN_RESPUESTA
    assert resultado.evidencias == []
    assert "respuesta_exacta" not in resultado.contrato_http


def test_deduplica_repeticiones_sin_fundir_secciones_distintas():
    lector = Lector()
    m = motor(candidatos=[candidato(TABLA, "a"), candidato(TABLA, "b"),
                          candidato(TABLA, "p2", seccion=["P2"])], lector=lector)
    resultado = m.consultar("ponderaciones", "s")
    assert resultado.diagnostico["duplicados"] == {"a": ["b"]}
    assert [e.chunk_id for e in resultado.evidencias] == ["a", "p2"]
    assert len(resultado.diagnostico["candidatos"]) == 3  # conserva procedencias para auditar


def test_variante_ortografica_respeta_presupuesto_y_pregunta_original(monkeypatch):
    from ipsis.lenguaje import ortografia

    monkeypatch.setattr(ortografia, "correcciones", lambda *_: {"pasword": "password"})
    lector = Lector()
    m = motor(lector=lector, idioma=object(), opciones={"recuperar_n": 3,
        "rerank_candidatos": 3, "rerank_k": 3, "leer_m": 2})
    pregunta = "¿Cómo cambio el pasword?"
    m.consultar(pregunta, "s")
    llamadas = m.retriever.consultas
    assert sum(limite for _, limite in llamadas) == 3
    assert len(llamadas) == 2
    assert all(q.pregunta_original == pregunta for q, _ in llamadas)
    assert "password" in llamadas[-1][0].contexto["pregunta_recuperacion"]
    assert lector.entradas[0][0] == pregunta


@pytest.mark.parametrize("version", [0, 1, 2, 3, 4, 5, 6, 7])
def test_migracion_sin_selector_elige_v2_y_es_idempotente(version):
    antigua = {"config_version": 5, "recuperacion": {"estrategia_chunking": "CH-F256", "coleccion_base": "x"}}
    antigua["config_version"] = version
    nueva = validar_motor(antigua)
    assert nueva["pipeline"] == "v2"
    assert nueva["config_version"] == 7
    assert validar_motor(nueva) == nueva
    assert "pipeline" not in antigua  # no modifica el archivo de entrada
    with pytest.raises(ValueError, match="pipeline"):
        validar_motor({**antigua, "pipeline": "v3"})


@pytest.mark.parametrize("version", [0, 5, 6, 7])
@pytest.mark.parametrize("pipeline", ["v1", "v2"])
def test_migracion_conserva_motor_elegido_explicitamente(version, pipeline):
    original = {"config_version": version, "pipeline": pipeline,
                "recuperacion": {"estrategia_chunking": "CH-F256", "coleccion_base": "x"}}
    assert validar_motor(original)["pipeline"] == pipeline
    assert migrar_motor(original)["pipeline"] == pipeline


def test_paquete_v2_no_importa_heuristicas_de_v1():
    raiz = Path(__file__).resolve().parents[2] / "ipsis/consulta/v2"
    prohibidos = ("ipsis.consulta.etapas", "ipsis.consulta.motor", "ipsis.consulta.seguimiento",
                  "ipsis.lenguaje.consultas", "ipsis.lenguaje.intenciones", "ipsis.lenguaje.aclaracion")
    for archivo in raiz.glob("*.py"):
        for nodo in ast.walk(ast.parse(archivo.read_text(encoding="utf-8"))):
            imports = [nodo.module] if isinstance(nodo, ast.ImportFrom) else (
                [a.name for a in nodo.names] if isinstance(nodo, ast.Import) else [])
            assert not any(m and m.startswith(prohibidos) for m in imports)


def test_reranker_ve_contexto_estructural_y_lector_solo_fuente():
    class Reranker:
        def puntuar(self, pregunta, pasajes):
            self.pasajes = pasajes
            return [1.0 for _ in pasajes]

    ch = candidato(TABLA)
    ch.chunk.texto_embedding = "Evaluación e instrumentos\n" + TABLA
    ch.chunk.metadata["structure_hash"] = "estructura"
    r, l = Reranker(), Lector("20%")
    m = motor(candidatos=[ch], reranker=r, lector=l)
    resultado = m.consultar("ponderaciones", "s")
    assert r.pasajes == [ch.chunk.texto_embedding]
    assert l.entradas[0][1] == [TABLA]
    assert resultado.evidencias[0].texto == TABLA
    assert "Evaluación e instrumentos" not in resultado.mensaje_usuario


def test_qa_en_otro_pasaje_no_desplaza_primero_ni_certifica_respuesta():
    correcto = candidato(TABLA, "primero")
    ajeno = candidato("La otra práctica tiene un arancel de $500.", "segundo", seccion=["P2"])
    resultado = motor(candidatos=[correcto, ajeno], lector=Lector("$500")).consultar(
        "¿Cuánto vale la primera práctica?", "s")
    assert [e.chunk_id for e in resultado.evidencias] == ["primero", "segundo"]
    assert resultado.estado == AssistantState.RELACIONADO
    assert resultado.contrato_http["estado"] == "relacionado"
    assert "respuesta_exacta" not in resultado.contrato_http
    assert "nivel_confianza" not in resultado.contrato_http
    assert resultado.diagnostico["lecturas"][1]["s_qa"] > .8
    assert resultado.diagnostico["decision"]["preferencia_qa"] == "segundo"
    assert resultado.diagnostico["decision"]["pasaje_principal"] == "primero"
    assert resultado.diagnostico["decision"]["qa_desplaza_pasajes"] is False


def test_presupuestos_v2_no_dependen_de_la_cascada_del_perfil():
    class CascadaProhibida:
        def __getattr__(self, nombre):
            raise AssertionError("v2 no debe leer presupuestos de v1")

    class Reranker:
        def puntuar(self, pregunta, pasajes):
            self.pasajes = pasajes
            return list(range(len(pasajes), 0, -1))

    r, l = Reranker(), Lector()
    m = motor(candidatos=[candidato(f"Pasaje {i}", str(i)) for i in range(12)],
        perfil=SimpleNamespace(perfil="medium", cascada=CascadaProhibida()), reranker=r, lector=l,
        opciones={"recuperar_n": 9, "rerank_candidatos": 7, "rerank_k": 5, "leer_m": 4,
                  "lote_reranker": 32})
    resultado = m.consultar("consulta", "s")
    assert m.retriever.consultas[0][1] == 9
    assert len(r.pasajes) == 7
    assert len(l.entradas[0][1]) == 4
    assert len(resultado.evidencias) == 3  # la publicación es independiente de las cuatro lecturas QA
    assert resultado.diagnostico["presupuestos"]["leer_m"] == 4


@pytest.mark.parametrize("opciones", [{"recuperar_n": 0}, {"leer_m": 101},
    {"rerank_candidatos": 61}, {"stride": 256, "max_tokens_pasaje": 128}, {"desconocida": True}])
def test_presupuestos_invalidos_fallan_antes_de_consultar(opciones):
    with pytest.raises(ValueError):
        motor(opciones=opciones)


def test_lector_por_campos_mapea_celdas_y_http_sin_inventar_span(tmp_path):
    from docx import Document
    from ipsis.ingesta.preparacion import preparar_docx
    from ipsis.ingesta.unidades import unidades_documento
    from ipsis.servicio.api import RespuestaConsulta, respuesta_publica

    doc = Document()
    doc.add_paragraph("Práctica I")
    t = doc.add_table(rows=4, cols=2)
    for i, valores in enumerate([("Instrumento", "Peso"), ("A", "20%"), ("B", "40%"), ("C", "40%")]):
        for j, valor in enumerate(valores):
            t.cell(i, j).text = valor
    archivo = tmp_path / "tabla.docx"
    doc.save(archivo)
    preparado = preparar_docx(archivo, doc_id="p1", doc_titulo="Reglamento")
    _, padres = unidades_documento(preparado)
    tabla = next(p for p in padres.values() if p.tipo_contenido == "tabla")
    c = RetrievalCandidate(chunk=tabla, signals=RetrievalSignals(fusion_rank=1, fusion_score=.03))
    lector = Lector("Peso: 20%")
    resultado = motor(candidatos=[c], lector=lector).consultar("ponderaciones", "s")
    assert "Peso: 20%" in lector.entradas[0][1][0]
    assert resultado.evidencias[0].texto == tabla.texto_fuente
    assert "A | 20%" in resultado.mensaje_usuario and "C | 40%" in resultado.mensaje_usuario
    bloque = next(b for b in resultado.contrato_http["bloques"] if b["tipo"] == "pasaje")
    refs = bloque["destacados"][0]["referencias"]
    assert [r["texto"] for r in refs] == ["Peso", "20%"]
    assert all(r["cell_id"] for r in refs)
    for ref in refs:
        assert preparado.texto_completo[ref["offset_inicio"]:ref["offset_fin"]] == ref["texto"]
    publica = RespuestaConsulta.model_validate(respuesta_publica(resultado.contrato_http))
    assert publica.status == "related" and publica.exact_answer is None
    assert len(publica.answer_blocks[-1].highlights[0].references) == 2
    assert resultado.diagnostico["senales"]["contextos"][0]["supuestos"]
    separador = motor(candidatos=[c], lector=Lector(":")).consultar("ponderaciones", "s")
    assert not any(b.get("destacados") for b in separador.contrato_http["bloques"])
    assert separador.evidencias[0].texto == tabla.texto_fuente


@pytest.mark.parametrize("sin_selector", [False, True])
def test_fabrica_default_v2_no_construye_cascada_y_se_puede_servir(monkeypatch, sin_selector):
    from fastapi.testclient import TestClient
    from ipsis.aplicacion.consulta import construir_motor
    from ipsis.configuracion.carga import cargar_proyecto, cargar_dominio
    from ipsis.nucleo.registro import RegistroComponentes
    from ipsis.servicio.api import crear_app
    from ipsis.lenguaje import intenciones

    def prohibido(*_):
        raise AssertionError("v2 no debe construir respaldo de intención")

    monkeypatch.setattr(intenciones, "construir_respaldo", prohibido)
    registro = RegistroComponentes()
    registro.registrar("retriever", "memoria_v2", lambda **_: Recuperador([candidato(TABLA)]))
    cfg = cargar_proyecto()
    assert cfg["pipeline"] == "v2"  # configuración materializada por ipsis init
    if sin_selector:
        cfg.pop("pipeline")
    cfg["componentes"]["retriever"] = "memoria_v2"
    dominio = {**cargar_dominio(), "perfil": "light"}
    m = construir_motor(cfg, cfg_dominio=dominio, registro=registro)
    assert isinstance(m, MotorV2)
    assert not hasattr(m, "politica_abstencion")
    assert m.abstencion_calibrada is False
    assert "intent_classifier" not in m.metadatos_auditoria["componentes"]
    assert m.metadatos_auditoria["componentes"]["answer_reader"] is None
    assert m.metadatos_auditoria["componentes_configurados"]["intent_classifier"] == "rules"
    with TestClient(crear_app(m)) as cliente:
        ready = cliente.get("/ready")
        assert ready.status_code == 200
        assert ready.json()["pipeline"] == "v2"
        assert cliente.get("/config").json()["response_mode"] == "ranked_sources"
        respuesta = cliente.post("/ask", json={"question": "ponderación", "session_id": "s"})
        assert respuesta.status_code == 200
        assert respuesta.json()["status"] == "related"
        assert respuesta.json()["response_mode"] == "ranked_sources"
        assert TABLA in respuesta.json()["answer"]
