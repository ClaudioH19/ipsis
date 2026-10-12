"""La preparación deja un resultado revisable y conserva el proyecto original."""
import json

from docx import Document
import pytest

from ipsis.aplicacion.preparacion import preparar_proyecto
from ipsis.configuracion.carga import cargar_proyecto, config_recuperacion
from ipsis.configuracion.archivos import volcar
from ipsis.ingesta.ingesta import huella_archivos, parsear_corpus


def proyecto(tmp_path):
    from ipsis.aplicacion.proyecto import inicializar_proyecto
    from ipsis.aplicacion.corpus import cargar_documentos
    import os

    origen = tmp_path / "original"
    origen.mkdir()
    previo = os.getcwd()
    try:
        os.chdir(origen)
        inicializar_proyecto(idioma="es", perfil="light")
        archivo = tmp_path / "doc.docx"
        d = Document()
        d.add_heading("Documento", 1)
        d.add_paragraph("<script>alert(1)</script>")
        d.add_table(rows=1, cols=2).cell(0, 0).text = "Valor"
        d.save(archivo)
        cargar_documentos([archivo])
    finally:
        os.chdir(previo)
    return origen


def test_prepara_copia_vista_escapada_y_config_sin_tocar_origen(tmp_path):
    origen = proyecto(tmp_path)
    antes = {p.relative_to(origen): p.read_bytes() for p in origen.rglob("*") if p.is_file()}
    r = preparar_proyecto(origen, tmp_path / "nuevo")
    assert r.indexado is False
    assert "&lt;script&gt;" in r.vista_previa.read_text(encoding="utf-8")
    assert "<script>" not in r.vista_previa.read_text(encoding="utf-8")
    assert json.loads(r.datos.read_text(encoding="utf-8"))["version"] == 1
    cfg = cargar_proyecto(str(r.proyecto / "config/engine.yaml"))
    assert cfg["pipeline"] == "v2"
    assert cfg["recuperacion"]["preparacion"]["modo"] == "estructural_v1"
    assert {p.relative_to(origen): p.read_bytes() for p in origen.rglob("*") if p.is_file()} == antes
    with pytest.raises(FileExistsError):
        preparar_proyecto(origen, r.proyecto)
    with pytest.raises(ValueError, match="fuera"):
        preparar_proyecto(origen, origen / "copia")


def test_huella_y_parseo_siguen_preparacion_activa(tmp_path, monkeypatch):
    origen = proyecto(tmp_path)
    r = preparar_proyecto(origen, tmp_path / "nuevo")
    monkeypatch.chdir(r.proyecto)
    cfg = cargar_proyecto()
    rec = config_recuperacion(cfg)
    assert rec["estrategia_chunking"] == "CH-STRUCT" and rec["con_header"] is True
    assert rec["coleccion_base"].endswith("_estructura1")
    assert rec["qdrant_path"].endswith("qdrant-estructura1")
    assert parsear_corpus("documents")[0].estructura is not None
    assert parsear_corpus("documents", "legacy")[0].estructura is None
    h = huella_archivos("documents", rec)
    cfg["recuperacion"]["preparacion"]["palabras_hijo"] = 90
    assert h != huella_archivos("documents", config_recuperacion(cfg))


def test_formato_no_soportado_falla_antes_de_crear_destino(tmp_path):
    from ipsis.configuracion.archivos import interpretar

    origen = proyecto(tmp_path)
    catalogo = origen / "documents/catalog.yaml"
    datos = interpretar(catalogo.read_text(encoding="utf-8"), "catalogo")
    datos["documentos"][0]["archivo"] = "nuevo.txt"
    catalogo.write_text(volcar(datos, "catalogo"), encoding="utf-8")
    (origen / "documents/nuevo.txt").write_text("Datos")
    destino = tmp_path / "rechazado"
    with pytest.raises(ValueError, match="DOCX"):
        preparar_proyecto(origen, destino)
    assert not destino.exists()
